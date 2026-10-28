from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path('output')
OUT.mkdir(exist_ok=True)
docx_path = OUT / 'WMUSD_Reserve_Confirmation_and_Attestation_Drafts.docx'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def borders(table, color='D9D9D9'):
    tbl = table._tbl; tblPr = tbl.tblPr
    b = tblPr.first_child_found_in('w:tblBorders')
    if b is None:
        b = OxmlElement('w:tblBorders'); tblPr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge; el = b.find(qn(tag))
        if el is None: el = OxmlElement(tag); b.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),color)

def add_field_table(doc, rows):
    table = doc.add_table(rows=0, cols=2); table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True; borders(table)
    for i,(a,b) in enumerate(rows):
        cells = table.add_row().cells
        cells[0].text=a; cells[1].text=b
        for c in cells: c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cells[0], 'EAF1FB'); shade(cells[1], 'FFFFFF')
        for run in cells[0].paragraphs[0].runs: run.bold=True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def heading(doc, text, level=1):
    p=doc.add_paragraph(style=f'Heading {level}'); p.add_run(text)
    return p

def body(doc, text, bold_lead=None):
    p=doc.add_paragraph()
    if bold_lead:
        r=p.add_run(bold_lead); r.bold=True
        p.add_run(text)
    else: p.add_run(text)
    p.paragraph_format.space_after=Pt(7); p.paragraph_format.line_spacing=1.12
    return p

doc=Document()
sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.85); sec.right_margin=Inches(.85)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(10.5); styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Aptos')
for s in ('Heading 1','Heading 2'):
    styles[s].font.name='Aptos Display'; styles[s].font.color.rgb=RGBColor(0,0,0)
styles['Heading 1'].font.size=Pt(16); styles['Heading 1'].font.bold=True
styles['Heading 2'].font.size=Pt(12); styles['Heading 2'].font.bold=True

p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('WMUSD Reserve Documentation Drafts'); r.font.name='Aptos Display'; r.font.size=Pt(22); r.bold=True; r.font.color.rgb=RGBColor(0,0,0)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Winmar Chain | Chain ID 12142816'); r.bold=True; r.font.color.rgb=RGBColor(31,78,121)
body(doc, 'This document contains working drafts for a bank confirmation letter and an independent reserve attestation concerning WMUSD, the proposed reserve-backed digital token of Winmar Holdings Pte. Ltd. The drafts must be completed and issued by the relevant bank and qualified attestor based on verified records.')

heading(doc,'WMUSD Reference Information',1)
add_field_table(doc,[
('Token name','Winmar USD Stablecoin'),('Token symbol','WMUSD'),('Issuing entity','Winmar Holdings Pte. Ltd.'),('Network','Winmar Chain'),('Chain ID','12142816'),('Reserve currency','United States Dollar (USD)'),('Proposed reserve amount','USD 100,000,000.00'),('Proposed maximum initial issuance','Up to 100,000,000 WMUSD, subject to verified eligible reserves and applicable law'),('Reserve standard','At least USD 1.00 of eligible reserve for each WMUSD in circulation'),
])

heading(doc,'Draft Bank Confirmation Letter',1)
body(doc,'[Print on official bank letterhead]')
add_field_table(doc,[('Date','[DD Month YYYY]'),('Recipient','[Independent Attestor / Audit Firm]'),('Subject','Confirmation of Funds Held for WMUSD Reserve Purposes')])
body(doc,'At the request and with the authorization of Winmar Holdings Pte. Ltd., we confirm that, according to our records, as of [date and time, including time zone], Winmar Holdings Pte. Ltd. maintained an available cash balance of USD 100,000,000.00 (United States Dollars One Hundred Million).')
body(doc,'The funds were held in an account maintained with [Bank Name], account ending [last four digits], in the name of Winmar Holdings Pte. Ltd.')
body(doc,'Based on the information and instructions provided to the Bank, the funds are intended to serve as reserve backing for WMUSD (Winmar USD Stablecoin), a proposed reserve-backed digital token issued by Winmar Holdings Pte. Ltd. on Winmar Chain, Chain ID 12142816.')
body(doc,'This confirmation is subject to the Bank’s records, applicable account terms, legal restrictions, liens, pledges, holds, and other encumbrances, if any.')
body(doc,'This letter confirms the account balance as of the stated date only. It does not constitute a bank guarantee, payment undertaking, redemption commitment, valuation, legal or regulatory opinion, audit of Winmar Holdings Pte. Ltd., audit of the WMUSD smart contract, or guarantee of the future value or redemption of WMUSD.')
body(doc,'This letter is issued solely for independent reserve verification and attestation purposes.')
body(doc,'For verification, the recipient may contact: [Authorized Bank Officer], [Title], [Official Bank Email], [Telephone], reference [Reference Number].')
body(doc,'Sincerely,\n\n[Signature]\n[Authorized Bank Officer]\n[Title]\n[Bank Name]')
body(doc,'If the Bank cannot certify that the funds are designated as reserve, use this alternative sentence: “The Bank confirms the account balance only. Any designation of the funds as reserve backing for WMUSD is the responsibility of Winmar Holdings Pte. Ltd. and does not constitute an undertaking or guarantee by the Bank.”', bold_lead='Alternative bank wording: ')

doc.add_page_break()
heading(doc,'Draft Independent WMUSD Reserve Attestation Report',1)
add_field_table(doc,[('To','The Board of Directors of Winmar Holdings Pte. Ltd. and WMUSD stakeholders'),('Subject','Verification of USD Reserve Backing for WMUSD'),('Reporting date','[DD Month YYYY]')])
body(doc,'We were engaged by Winmar Holdings Pte. Ltd. to perform procedures over the USD reserve intended to support the issuance of WMUSD (Winmar USD Stablecoin) on Winmar Chain, Chain ID 12142816.')
body(doc,'Management of Winmar Holdings Pte. Ltd. is responsible for maintaining the USD reserve, providing complete and accurate supporting documents, ensuring that reserve assets are properly identified, preventing double-counting or pledging of reserve assets, and determining the amount of WMUSD issued and in circulation.')
heading(doc,'Applicable Reserve Criterion',2)
body(doc,'Each WMUSD in circulation should be backed by not less than one United States Dollar held as an eligible reserve asset. The reserve amount and WMUSD circulating supply must be measured as of the same reporting date and time.')
heading(doc,'Reported Reserve Position',2)
add_field_table(doc,[('Eligible USD reserve balance','USD 100,000,000.00'),('WMUSD circulating supply','[Amount as of reporting date]'),('Maximum issuance supported by verified reserve','[Amount] WMUSD'),('Reserve coverage ratio','[Percentage]'),('Reserve status','[Sufficient / Insufficient / Subject to stated limitations]')])
body(doc,'Based on the documents provided and the bank confirmation received directly from [Bank Name], the verified eligible USD reserve was USD 100,000,000.00 as of [date and time].')
body(doc,'This report is limited to the reserve position at the stated date. It does not guarantee future reserve levels, future redemption, token price, legal compliance, solvency, or the performance of the WMUSD smart contract. It should not be interpreted as a bank guarantee or as a guarantee by the attestor of the future value of WMUSD.')
heading(doc,'Agreed-Upon Procedures Version',2)
body(doc,'If the engagement is an Agreed-Upon Procedures engagement, replace the conclusion above with: “We report the factual findings from the procedures performed. We do not express an audit opinion or assurance conclusion.”')
body(doc,'Attestor: [Name of Attestation Firm] | [License / Registration Number] | [Address] | [Authorized Signatory] | [Professional Title] | [Date] | [Engagement Reference]')

heading(doc,'Completion Notes',1)
body(doc,'The bank should send its confirmation directly to the attestor through a verifiable official channel. Do not publish a full account number, private banking information, or an unverified reserve claim. If no WMUSD has been issued yet, report “WMUSD circulating supply: 0” and describe USD 100,000,000.00 as verified reserve capacity rather than current circulating-token backing.')

doc.save(docx_path)
print(docx_path)
